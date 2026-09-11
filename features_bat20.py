"""
battery20_features.py — TWENTY physically-motivated BATTERY-storage main features, computed only
from the AEW smart-meter CSVs (no weather, no external data).

This is a fork of pv20_features.py: it reuses the SAME file / store reading layer (compatible with
the SAME aew_store built by build_store.py), but replaces the analysis with features meant to detect
a BATTERY behind the meter — whether charging (from the grid or from PV) or discharging (to the house
or to the grid) — rather than PV alone.

Run (inside the Renku session, from ~/work/mechenergy):
    FAST PATH — after a one-time  python build_store.py ../aew-data/test-blob/input_data -o aew_store :
        python battery20_features.py --store aew_store --mp-ids 53628,45390 -o two_houses.csv
        python battery20_features.py --store aew_store -o all_houses.csv
        (if ./aew_store exists it is used automatically, even when a CSV root is given)
    CSV PATH — no store yet:
        python battery20_features.py ../aew-data/test-blob/input_data --max-houses 10 -o ten_houses.csv
        python battery20_features.py ../aew-data/test-blob/input_data --mp-ids 53628,45390,49688 -o some.csv
        python battery20_features.py ../aew-data/test-blob/input_data --mp-ids-file labelled_ids.txt -o labelled.csv

HOW THE FILES ARE READ — unchanged from pv20_features.py: with a house list, each monthly file is
grep-filtered as TEXT first and only those lines are parsed; without one, the file is streamed in
chunks and reduced at once, so the 96-column matrix never sits in memory.

Input files: only files named LG_AIM2Hackerdays_kWh_*.csv are read (change with --pattern).
Layout of a data row (';'-separated):  MP ID ; meter id ; OBIS-Code ; Datum ; PLZ ; 00:15 ; ... ; 24:00
    OBIS 1-1:1.29.0*255 = energy IMPORTED from the grid  (kWh per 15 min)
    OBIS 1-1:2.29.0*255 = energy EXPORTED to the grid    (kWh per 15 min)
THE DATA ARE NON-NEGATIVE (two separate counters); negative raw values are floored at 0 and
counted in n_negative_raw.

==============================================================================
WHY THESE FEATURES POINT TO A BATTERY AND NOT TO PV
==============================================================================
PV output is a smooth function of the sun's position: it exists only between sunrise and sunset,
is unimodal (peaks near solar noon), varies smoothly with cloud cover, and is strongly seasonal
(much less in winter).  A battery breaks every one of these properties:
  - it can export, or stop importing, at ANY hour, including hours with no possible sunlight
    (this is the single cleanest "not the sun" signal used below: the night window);
  - a battery inverter / charger often runs at a near-constant rate for a sustained stretch
    (SOC- or tariff-controlled), producing a FLAT plateau, unlike PV's continuously varying bell;
  - a battery can cycle several times a day (charge off-peak at night, discharge at the evening
    peak, charge again from midday PV surplus, ...), producing several on/off "bursts" per day
    where PV alone gives one continuous episode from sunrise to sunset;
  - a battery stores energy across the day, so its effect on the IMPORT shape is much less
    seasonal than PV's (PV empties the midday only in summer; a battery flattens the day all year);
  - battery self-consumption shifts energy in time: PV surplus exported / stored around midday
    shows up as REDUCED import later the same day (typically the evening peak), a same-day
    cross-channel correlation that PV alone cannot produce (PV-only self-consumption without
    storage cannot make the evening cheaper using energy captured at midday).

CAVEATS (read these before trusting a house's score):
  - Off-peak GRID CHARGING (features 11-13, 20) is close to indistinguishable from a home EV
    charger on a night tariff.  A high score there is evidence of "something drawing power from
    the grid on a schedule", not proof of a stationary battery; check it against the export
    evidence (features 1-9) or house metadata if you need to tell them apart.
  - export_summer_winter_flatness (9) and import_seasonal_flatness (17) are computed as
    -abs(log(ratio)), so a value near 0 means "steady across seasons" (battery-like) and a very
    negative value means "strongly seasonal" (PV-like, or an ordinary heating load).  BUT a house
    with no export at all also gives a ratio eps/eps = 1 -> flatness = 0, which looks
    "battery-like" for the wrong reason (absence of data, not evidence of a battery).  Always
    read (9) together with has_export_register / has_real_export, and (17) together with
    n_summer_days / n_winter_days, before interpreting a value near 0 as real evidence.
  - self_consumption_shift_score (20) needs enough paired days with real values on both sides;
    it is NaN -> filled with 0 (neutral) when a house has too few overlapping days.
  - None of this is a certified detector: it surfaces evidence consistent with a battery so that
    a human or a downstream classifier can rank and inspect houses, not a certified detection.

==============================================================================
THE TWENTY FEATURES   (summer = May-Aug, winter = Nov-Feb, "night" = 22h-06h)
==============================================================================
A. EXPORT COUNTER — timing and shape evidence of a battery that PV alone cannot produce
 1. export_kwh_total              [kWh, >= 0]  total energy ever exported (0 = nothing behind the
    meter that can push energy out).  Baseline size metric, shared with any DER.
 2. export_days_frac              [0..1]  share of days with > 0.1 kWh exported.
 3. export_offsolar_over_import   [>= 0]  export in the night window / import, all days.  The sun
    cannot do this at all; > 0 here is direct evidence of storage (or another non-solar generator).
 4. export_offsolar_days_frac     [0..1]  share of days with any measurable (> ~0.02 kWh) export in
    that same night window: persistence of the signal, not just its size.
 5. export_flat_plateau_frac      [0..1]  of the quarter-hours actively exporting, the share within
    10 % of that day's maximum, averaged over exporting days.  PV's bell curve spends little time
    that close to its peak; a rate-limited battery inverter / discharge sits there for a long stretch.
 6. export_bursts_per_day         [>= 0]  mean number of separate export episodes per day (rising
    edges through the noise floor).  PV alone -> ~1 (sunrise to sunset); a battery cycling more
    than once a day -> 2 or more.
 7. export_max_kw                 [kW]  95th percentile of the daily maximum export power ~ size of
    the inverter (0 if none).
 8. export_max_over_mean_import   [>= 0]  that maximum relative to the house's mean import power.
 9. export_summer_winter_flatness [<= 0]  -abs(log((summer export/day + eps) / (winter export/day
    + eps))).  Near 0 = export is similar in both seasons (a battery buffers all year); very
    negative = sharply seasonal (PV, which barely exports in winter).  See the caveat above.

B. IMPORT COUNTER — charge / discharge shape evidence (works on every meter)
10. import_zero_hours_night       [h, 0..8]  hours per day (22h-06h) with import <= 20 W, all year.
    The sun cannot cause this; zero import at night means the house is being fed from elsewhere:
    a battery (or it is empty).
11. import_night_charge_kw_over_mean  [>= 0]  95th percentile, over the house's days, of its mean
    import power in the 23h-06h window, divided by its overall mean import power.  A sustained high
    level here (well above 1) is consistent with scheduled off-peak grid charging.  See the EV
    charger caveat above.
12. import_night_charge_flatness  [>= 0]  coefficient of variation (std / mean) of import power
    inside that same 23h-06h window, averaged over nights with enough valid data.  Lower = steadier
    (a charger holding a fixed rate) than normal, noisier household night load.
13. import_night_charge_persistence  [0..1]  share of nights on which the mean 23h-06h import
    exceeds 1.5x the house's overall mean import: how often the elevated night pattern recurs,
    complementing the magnitude of (11) with frequency.
14. import_evening_peak_suppression  [<= 1]  1 - (mean import 17h-21h) / (mean import over the rest
    of the day), averaged over days.  Positive and large = the evening peak is unusually low relative
    to the rest of the day's draw, consistent with peak-shaving discharge.
15. import_ramp_events_night      [>= 0]  mean count, per day, of quarter-hour-to-quarter-hour import
    jumps larger than RAMP_KW within 22h-06h.  Captures the on/off switching of a charger or a
    discharge controller, as opposed to the smoother drift of normal appliance cycling.
16. import_zero_hours_midday_allyear  [h, 0..8]  hours per day (9h-17h) with import <= 20 W,
    computed over ALL months, not just summer.  If this persists in winter (weak sun), it is not
    explained by PV alone and is more consistent with a battery buffering the midday load all year.
17. import_seasonal_flatness      [<= 0]  same -abs(log(ratio)) construction as (9), but on the
    house's total daily import (summer vs winter).  Near 0 = small seasonal shape difference
    (a buffered load); very negative = sharply seasonal (e.g. heating, or PV emptying only the
    midday hours in summer).  See the caveat above.
18. import_midday_night_ratio_allyear  [>= 0]  mean import 11h-15h / max(mean import 22h-06h, 20 W),
    computed over ALL months.  General consumption-shape descriptor; a battery discharging at night
    lowers this ratio even outside summer.

C. JOINT — evidence spanning both counters on the same day
19. daily_activity_cycles         [>= 0]  mean per day of (export bursts, feature 6) plus (distinct
    23h-06h charging episodes above CHARGE_KW).  A system cycling through several charge / discharge
    events per day scores higher here than one continuous PV episode.
20. self_consumption_shift_score  [-1..1]  the NEGATIVE of the day-to-day Pearson correlation, per
    house, between that day's midday (11h-15h) export and the SAME day's evening-peak (17h-21h)
    import.  A positive score means: on days when more was exported at midday, less was imported
    that same evening — the time-shift signature of storing midday surplus and using it later,
    which plain PV self-consumption (no storage) cannot produce on its own.

Export noise floor: EXPORT_NOISE_KW = 0.05 kW (meter creep).  has_real_export = 1 when the house's
typical daily maximum export reaches EXPORT_REAL_KW (0.1 kW).  has_offsolar_export = 1 when the
house shows real export in the night window on a non-trivial share of days.

Data-quality columns: n_days, n_summer_days, n_winter_days, has_export_register, has_real_export,
has_offsolar_export, export_max_raw_kw, mean_import_kw, n_negative_raw.
"""
from __future__ import annotations

import argparse
import glob
import hashlib
import io
import os
import re
import shutil
import subprocess
import tempfile
import time
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from typing import Dict, Iterable, Iterator, List, Optional, Set, Tuple

import numpy as np
import pandas as pd

N_SLOTS = 96
SUMMER_MONTHS = (5, 6, 7, 8)
WINTER_MONTHS = (11, 12, 1, 2)
ZERO_KW = 0.02            # import <= 20 W == "the house imports nothing"
EXPORT_NOISE_KW = 0.05    # export below 50 W is meter noise, not generation
EXPORT_REAL_KW = 0.10     # a house "really exports" if its typical daily max export reaches 100 W
CHARGE_KW = 0.50          # "significant charging" power used for burst / persistence detection
RAMP_KW = 0.30            # a step >= 300 W between consecutive quarter-hours counts as a "ramp event"
OFFSOLAR_KWH_DAY = 0.02   # export threshold (per day, in the night window) to count as a "hit" for feature 4

SLOT_HOUR = (np.arange(N_SLOTS) + 0.5) / 4.0


def _mask(a: float, b: float) -> np.ndarray:
    """Quarter-hours whose centre lies in [a, b) — wraps past midnight if a > b."""
    return (SLOT_HOUR >= a) & (SLOT_HOUR < b) if a < b else (SLOT_HOUR >= a) | (SLOT_HOUR < b)


MID_MASK = _mask(11, 15)      # midday centred on solar noon
PEAK_MASK = _mask(17, 21)     # evening demand peak
NIGHT_MASK = _mask(22, 6)     # 22h-06h: the sun cannot explain zero import or any export here
CHARGE_MASK = _mask(23, 6)    # off-peak charging window (adjust to the local tariff if needed)
DAY_MASK = _mask(9, 17)       # daylight hours, used for the all-year zero-import check

FEATURE_NAMES = [
    "export_kwh_total", "export_days_frac", "export_offsolar_over_import",
    "export_offsolar_days_frac", "export_flat_plateau_frac", "export_bursts_per_day",
    "export_max_kw", "export_max_over_mean_import", "export_summer_winter_flatness",
    "import_zero_hours_night", "import_night_charge_kw_over_mean", "import_night_charge_flatness",
    "import_night_charge_persistence",
    "import_evening_peak_suppression", "import_ramp_events_night",
    "import_zero_hours_midday_allyear", "import_seasonal_flatness", "import_midday_night_ratio_allyear",
    "daily_activity_cycles", "self_consumption_shift_score",
]

VCOLS = [f"v{k}" for k in range(N_SLOTS)]
KEY = ["mp_id", "channel", "date"]

# ---- the meter files (layout handling identical to pv20_features.py) ----
FILE_GLOB = "LG_AIM2Hackerdays_kWh_*.csv"
OBIS_IMPORT = "1-1:1.29.0*255"
OBIS_EXPORT = "1-1:2.29.0*255"
OBIS_RE = re.compile(r"^\d+-\d+:\d+\.\d+\.\d+")
DATE_RE = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")
METER_RE = re.compile(r"^[A-Z]{2}\d{6,}")


def file_layout(path: str) -> dict:
    """{'n_lead': leading columns, 'roles': {'mp_id': i, 'obis': i, 'date': i, 'meter': i|None, 'plz': i|None}}"""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        f.readline()
        first = f.readline().rstrip("\r\n").split(";")
    fields = first[:-1] if first and first[-1].strip() == "" else first
    n_lead = len(fields) - N_SLOTS
    if n_lead < 3:
        raise ValueError(f"{path}: expected >= 3 leading columns + {N_SLOTS} values, got {len(fields)} fields")
    lead = [x.strip() for x in fields[:n_lead]]
    roles: dict = {"meter": None, "plz": None}
    for i, v in enumerate(lead):
        if "obis" not in roles and OBIS_RE.match(v):
            roles["obis"] = i
        elif "date" not in roles and DATE_RE.match(v):
            roles["date"] = i
        elif roles["meter"] is None and METER_RE.match(v):
            roles["meter"] = i
    if "obis" not in roles or "date" not in roles:
        raise ValueError(f"{path}: could not recognise OBIS / date in the first data row {lead}")
    rest = [i for i in range(n_lead) if i not in (roles["obis"], roles["date"], roles["meter"])]
    roles["mp_id"] = rest[0]
    roles["plz"] = rest[1] if len(rest) > 1 else None
    return {"n_lead": n_lead, "roles": roles}


def layout_names(lay: dict) -> List[str]:
    inv = {i: r for r, i in lay["roles"].items() if i is not None}
    return [inv.get(i, f"x{i}") for i in range(lay["n_lead"])] + VCOLS


# =============================================================================
# 1. reading the CSVs (unchanged from pv20_features.py)
# =============================================================================
def find_files(root: str, month: Optional[str] = None, max_files: Optional[int] = None,
               pattern: str = FILE_GLOB) -> List[str]:
    files = sorted(glob.glob(os.path.join(root, "**", pattern), recursive=True))
    if month:
        files = [f for f in files if month.lower() in f.lower()]
    if max_files:
        files = files[:max_files]
    if not files:
        raise SystemExit(f"no files named {pattern} found below {root}" + (f" matching '{month}'" if month else ""))
    return files


def first_mp_ids(path: str, n: int, max_lines: int = 200_000) -> List[str]:
    ids: List[str] = []
    seen: Set[str] = set()
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        f.readline()
        for i, line in enumerate(f):
            mp = line.split(";", 1)[0].strip()
            if mp and mp not in seen:
                seen.add(mp); ids.append(mp)
                if len(ids) >= n:
                    break
            if i >= max_lines:
                break
    return ids


def filter_lines(path: str, mp_ids: Iterable[str], cache_dir: Optional[str] = None) -> str:
    ids = sorted(set(str(m) for m in mp_ids))
    if not ids:
        return ""
    cache_file = None
    if cache_dir:
        st = os.stat(path)
        key = hashlib.md5(f"{os.path.abspath(path)}|{st.st_size}|{int(st.st_mtime)}|{','.join(ids)}".encode()).hexdigest()
        cache_file = os.path.join(cache_dir, f"{key}.txt")
        if os.path.exists(cache_file):
            with open(cache_file, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
    text = None
    if shutil.which("grep"):
        env = {**os.environ, "LC_ALL": "C"}
        if len(ids) <= 50:
            pattern = "^(" + "|".join(re.escape(m) for m in ids) + ");"
            r = subprocess.run(["grep", "-E", pattern, path], capture_output=True, text=True,
                               encoding="utf-8", errors="replace", env=env)
            if r.returncode in (0, 1):
                text = r.stdout
        else:
            with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False) as tf:
                tf.write("".join(f"{m};\n" for m in ids))
            try:
                r = subprocess.run(["grep", "-F", "-f", tf.name, path], capture_output=True, text=True,
                                   encoding="utf-8", errors="replace", env=env)
            finally:
                os.unlink(tf.name)
            if r.returncode in (0, 1):
                keep = set(ids)
                text = "".join(l + "\n" for l in r.stdout.splitlines() if l.split(";", 1)[0].strip() in keep)
    if text is None:
        keep = set(ids)
        out = []
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                if line.split(";", 1)[0].strip() in keep:
                    out.append(line)
        text = "".join(out)
    if cache_file:
        os.makedirs(cache_dir, exist_ok=True)
        with open(cache_file, "w", encoding="utf-8") as f:
            f.write(text)
    return text


def _parse(source, chunksize: Optional[int], skip_header: bool, lay: dict):
    names = layout_names(lay)
    dtypes = {c: "string" for c in names[:lay["n_lead"]]}
    dtypes.update({v: "float32" for v in VCOLS})
    return pd.read_csv(source, sep=";", header=None, skiprows=1 if skip_header else 0, names=names, index_col=False,
                       dtype=dtypes, na_values=["", " ", "-"], keep_default_na=False, chunksize=chunksize,
                       encoding="utf-8", encoding_errors="replace", engine="c")


def _post(lab: pd.DataFrame, vals: np.ndarray):
    lab = lab.reset_index(drop=True)
    lab["n_neg"] = (vals < 0).sum(axis=1)
    np.maximum(vals, np.float32(0.0), out=vals)
    is_exp = (lab["channel"] == "export").values
    lab["vmax_raw"] = np.nanmax(np.where(np.isfinite(vals), vals, -1.0), axis=1)
    if is_exp.any():
        ve = vals[is_exp]
        ve[ve < EXPORT_NOISE_KW] = 0.0
        vals[is_exp] = ve
    return lab, vals


def _finish(ch: pd.DataFrame):
    obis = ch["obis"].astype(str).str.strip()
    lab = pd.DataFrame({
        "mp_id": ch["mp_id"].astype(str).str.strip(),
        "channel": np.where(obis == OBIS_IMPORT, "import", np.where(obis == OBIS_EXPORT, "export", "other")),
        "date": pd.to_datetime(ch["date"].astype(str).str.strip(), format="%d.%m.%Y", errors="coerce"),
    }, index=ch.index)
    sel = (lab["channel"] != "other").values & lab["date"].notna().values
    if not sel.any():
        return None
    vals = ch.loc[sel, VCOLS].to_numpy(dtype="float32") * np.float32(4.0)     # kWh / 15 min -> kW
    return _post(lab[sel].copy(), vals)


def iter_aew_csv(path: str, mp_ids: Optional[Iterable[str]] = None,
                 chunksize: int = 200_000, cache_dir: Optional[str] = None) -> Iterator[Tuple[pd.DataFrame, np.ndarray]]:
    lay = file_layout(path)
    if mp_ids is not None:
        text = filter_lines(path, mp_ids, cache_dir)
        if not text.strip():
            return
        out = _finish(_parse(io.StringIO(text), None, skip_header=False, lay=lay))
        if out is not None:
            yield out
        return
    for ch in _parse(path, chunksize, skip_header=True, lay=lay):
        out = _finish(ch)
        if out is not None:
            yield out


# =============================================================================
# 1b. reading the binary store built by build_store.py (unchanged from pv20_features.py)
# =============================================================================
STORE_DEFAULT = "aew_store"
N_BUCKETS = 256
GERMAN_MONTHS = {m: i + 1 for i, m in enumerate(
    ["januar", "februar", "märz", "april", "mai", "juni", "juli", "august", "september", "oktober", "november", "dezember"])}


def bucket_of(mp_id: str) -> int:
    try:
        return int(str(mp_id).strip()) % N_BUCKETS
    except ValueError:
        return 0


def _month_to_dates(month: Optional[str]):
    if not month:
        return None
    t = month.strip().lower()
    yr = re.search(r"(20\d{2})", t)
    if not yr:
        print(f"warning: --month '{month}' has no year -> ignored for the store", flush=True)
        return None
    y = int(yr.group(1))
    mo = next((v for k, v in GERMAN_MONTHS.items() if k in t), None)
    if mo is None:
        return pd.Timestamp(y, 1, 1).date(), pd.Timestamp(y + 1, 1, 1).date()
    end = pd.Timestamp(y + (mo == 12), mo % 12 + 1, 1).date()
    return pd.Timestamp(y, mo, 1).date(), end


def store_buckets(store: str) -> List[str]:
    dirs = sorted(glob.glob(os.path.join(store, "bucket=*")), key=lambda d: int(d.rsplit("=", 1)[1]))
    if not dirs:
        raise SystemExit(f"{store} is not a store (no bucket=* folders); build it with build_store.py")
    return dirs


def read_store(store: str, mp_ids: Optional[Iterable[str]] = None, buckets: Optional[Iterable[str]] = None,
               month: Optional[str] = None):
    import pyarrow.compute as pc  # noqa: F401
    import pyarrow.dataset as ds

    dates = _month_to_dates(month)
    if mp_ids is not None:
        ids = sorted(set(str(m) for m in mp_ids))
        by_bucket: Dict[int, List[str]] = {}
        for m in ids:
            by_bucket.setdefault(bucket_of(m), []).append(m)
        jobs = [(os.path.join(store, f"bucket={b}"), ids_b) for b, ids_b in sorted(by_bucket.items())]
    else:
        jobs = [(d, None) for d in (buckets if buckets is not None else store_buckets(store))]
    for d, ids_b in jobs:
        if not os.path.isdir(d):
            continue
        flt = None
        if ids_b is not None:
            flt = ds.field("mp_id").isin(ids_b)
        if dates is not None:
            f2 = (ds.field("date") >= dates[0]) & (ds.field("date") < dates[1])
            flt = f2 if flt is None else (flt & f2)
        tab = ds.dataset(d, format="parquet").to_table(filter=flt, columns=["mp_id", "channel", "date"] + VCOLS,
                                                       use_threads=False)
        if tab.num_rows == 0:
            continue
        lab = pd.DataFrame({
            "mp_id": tab.column("mp_id").to_pandas().astype(str),
            "channel": tab.column("channel").to_pandas().astype(str),
            "date": pd.to_datetime(tab.column("date").to_pandas()),
        })
        vals = np.array(tab.select(VCOLS).to_pandas().to_numpy(dtype="float32"), copy=True)
        yield _post(lab, vals)


def first_store_ids(store: str, n: int) -> List[str]:
    import pyarrow.parquet as pq
    ids: List[str] = []
    for d in store_buckets(store):
        col = pq.read_table(os.path.join(d, "part.parquet"), columns=["mp_id"]).column("mp_id").to_pylist()
        ids.extend(sorted(set(col)))
        if len(ids) >= n:
            break
    return ids[:n]


# =============================================================================
# 2. reduce every (meter, channel, day) row of 96 values to per-day battery numbers
# =============================================================================
_F = lambda m: m.astype(np.float32)
_W_MID, _W_PEAK, _W_NIGHT, _W_DAY, _W_CHARGE = _F(MID_MASK), _F(PEAK_MASK), _F(NIGHT_MASK), _F(DAY_MASK), _F(CHARGE_MASK)
_W_ONE = np.ones(N_SLOTS, dtype=np.float32)


def _bursts(active: np.ndarray) -> np.ndarray:
    """Count rising edges (False -> True) per row of a boolean (n, 96) activity matrix."""
    n = active.shape[0]
    padded = np.concatenate([np.zeros((n, 1), dtype=bool), active], axis=1)
    rising = (~padded[:, :-1]) & padded[:, 1:]
    return rising.sum(axis=1).astype(np.float32)


def reduce_days(lab: pd.DataFrame, vals: np.ndarray) -> pd.DataFrame:
    """Per-day numbers the twenty battery features are built from (values are kW, >= 0)."""
    finite = np.isfinite(vals)
    v = np.where(finite, vals, np.float32(0.0)).astype(np.float32, copy=False)
    fin = finite.astype(np.float32)
    d = lab[KEY + ["n_neg", "vmax_raw"]].copy() if "vmax_raw" in lab else lab[KEY + ["n_neg"]].assign(vmax_raw=np.nan)

    with np.errstate(invalid="ignore", divide="ignore"):
        d["tot"] = v @ _W_ONE
        d["mid_tot"] = v @ _W_MID
        d["peak_tot"] = v @ _W_PEAK
        d["night_tot"] = v @ _W_NIGHT          # the 22h-06h "cannot be the sun" window
        d["day_tot"] = v @ _W_DAY

        d["n_valid"] = fin @ _W_ONE
        d["n_valid_mid"] = fin @ _W_MID
        d["n_valid_peak"] = fin @ _W_PEAK
        d["n_valid_night"] = fin @ _W_NIGHT
        d["n_valid_day"] = fin @ _W_DAY

        vmax = np.where(finite, vals, np.float32(-1.0)).max(axis=1)
        d["vmax"] = np.where(vmax < 0, np.nan, vmax)

        zero = ((v <= ZERO_KW) & finite).astype(np.float32)
        d["zero_q_night"] = zero @ _W_NIGHT
        d["zero_q_day"] = zero @ _W_DAY

        # off-peak charging window (23h-06h): level + steadiness, via sufficient statistics
        d["charge_sum"] = v @ _W_CHARGE
        d["charge_sqsum"] = (v * v) @ _W_CHARGE
        d["charge_cnt"] = fin @ _W_CHARGE

        # export activity: bursts (on/off cycles) and how long is spent near the day's max
        active = (v > EXPORT_NOISE_KW) & finite
        d["bursts"] = _bursts(active)
        dmax = np.where(active.any(axis=1), v.max(axis=1), np.nan)
        near_max = active & (v >= 0.9 * dmax[:, None])
        n_active = active.sum(axis=1)
        d["plateau_frac"] = np.where(n_active > 0, near_max.sum(axis=1) / np.maximum(n_active, 1), np.nan)

        # night ramp events: level changes larger than RAMP_KW within 22h-06h
        diffs = np.abs(np.diff(v, axis=1))
        both_night = NIGHT_MASK[:-1] & NIGHT_MASK[1:]
        d["ramp_events_night"] = ((diffs > RAMP_KW) & both_night[None, :]).sum(axis=1).astype(np.float32)

        # distinct charging episodes within 23h-06h (for the joint cycling feature)
        charge_active = (v > CHARGE_KW) & CHARGE_MASK[None, :] & finite
        d["charge_bursts"] = _bursts(charge_active)

    return d


def load_days(root: Optional[str] = None, verbose: bool = True, month: Optional[str] = None,
              max_files: Optional[int] = None, max_houses: Optional[int] = None,
              mp_ids: Optional[Iterable[str]] = None, pattern: str = FILE_GLOB, workers: int = 4,
              cache_dir: Optional[str] = None, store: Optional[str] = None,
              use_days_cache: bool = True) -> pd.DataFrame:
    if store:
        return load_days_from_store(store, verbose, month, max_houses, mp_ids, workers, use_days_cache)
    if not root:
        raise SystemExit("give the CSV root folder, or --store aew_store (build it with build_store.py)")
    files = find_files(root, month, max_files, pattern)
    keep: Optional[Set[str]] = set(str(m) for m in mp_ids) if mp_ids is not None else None
    if keep is None and max_houses:
        keep = set(first_mp_ids(files[0], max_houses))
        if verbose:
            print(f"first {len(keep)} MP IDs of {os.path.basename(files[0])}: {sorted(keep)}")

    def one_file(f: str) -> Tuple[str, List[pd.DataFrame], int, Set[str]]:
        parts, n_rows, ids = [], 0, set()
        for lab, vals in iter_aew_csv(f, mp_ids=keep, cache_dir=cache_dir if keep is not None else None):
            parts.append(reduce_days(lab, vals))
            n_rows += len(lab); ids.update(lab["mp_id"].unique())
        return f, parts, n_rows, ids

    parts: List[pd.DataFrame] = []
    n_workers = max(1, workers) if keep is not None else 1
    with ThreadPoolExecutor(max_workers=n_workers) as pool:
        for i, (f, fparts, n_rows, ids) in enumerate(pool.map(one_file, files)):
            if verbose:
                print(f"[{i + 1}/{len(files)}] {os.path.relpath(f, root)} ... {n_rows:,} rows, {len(ids):,} houses", flush=True)
            parts.extend(fparts)
    days = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame(columns=KEY + ["n_neg"])
    return _fix_multimeter(days, verbose, lambda ids: [_exact_days_for(f, ids, cache_dir) for f in files])


def load_days_from_store(store: str, verbose: bool = True, month: Optional[str] = None,
                         max_houses: Optional[int] = None, mp_ids: Optional[Iterable[str]] = None,
                         workers: int = 4, use_days_cache: bool = True) -> pd.DataFrame:
    keep: Optional[List[str]] = sorted(set(str(m) for m in mp_ids)) if mp_ids is not None else None
    if keep is None and max_houses:
        keep = first_store_ids(store, max_houses)
        if verbose:
            shown = keep if len(keep) <= 20 else keep[:10] + ["..."] + keep[-5:]
            print(f"first {len(keep)} MP IDs of the store: {shown}")
    t0 = time.time()
    days_cache = os.path.join(store, "_battery_days.parquet")
    if keep is None and not month and use_days_cache and os.path.exists(days_cache):
        days = pd.read_parquet(days_cache)
        if verbose:
            print(f"store: battery per-day table for all houses loaded from {days_cache} ({len(days):,} rows, "
                  f"{time.time() - t0:.1f} s)  [--refresh-days to recompute]", flush=True)
        return days
    if keep is not None:
        by_bucket: Dict[int, List[str]] = {}
        for m in keep:
            by_bucket.setdefault(bucket_of(m), []).append(m)
        jobs = [(store, os.path.join(store, f"bucket={b}"), month, ids) for b, ids in sorted(by_bucket.items())]
        if len(jobs) <= 4 or workers <= 1:
            parts = [p for job in jobs for p in _reduce_bucket(job)]
        else:
            parts = []
            with ProcessPoolExecutor(max_workers=max(1, min(workers, os.cpu_count() or 1)),
                                     initializer=_single_thread_blas) as pool:
                for i, bparts in enumerate(pool.map(_reduce_bucket, jobs)):
                    parts.extend(bparts)
                    if verbose and (i + 1) % 16 == 0:
                        print(f"  {i + 1}/{len(jobs)} buckets ({time.time() - t0:.0f} s)", flush=True)
        if verbose:
            n = sum(len(p) for p in parts)
            print(f"store: {n:,} meter-days for {len(keep):,} houses in {time.time() - t0:.1f} s", flush=True)
    else:
        buckets = store_buckets(store)
        parts = []
        with ProcessPoolExecutor(max_workers=max(1, min(workers, os.cpu_count() or 1)),
                                 initializer=_single_thread_blas) as pool:
            for i, bparts in enumerate(pool.map(_reduce_bucket, [(store, d, month) for d in buckets])):
                parts.extend(bparts)
                if verbose and (i + 1) % 32 == 0:
                    print(f"  {i + 1}/{len(buckets)} buckets ({time.time() - t0:.0f} s)", flush=True)
        if verbose:
            print(f"store: all {len(buckets)} buckets reduced in {time.time() - t0:.0f} s", flush=True)
    days = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame(columns=KEY + ["n_neg"])
    days = _fix_multimeter(days, verbose, lambda ids: [_exact_days_from_store(store, ids, month)])
    if keep is None and not month and use_days_cache:
        days.to_parquet(days_cache, index=False)
        if verbose:
            print(f"store: battery per-day table saved to {days_cache} (next all-houses run skips the reduction)", flush=True)
    return days


def _single_thread_blas() -> None:
    for k in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS"):
        os.environ[k] = "1"


def _reduce_bucket(args) -> List[pd.DataFrame]:
    store, d, month = args[:3]
    ids = args[3] if len(args) > 3 else None
    if ids is None:
        return [reduce_days(lab, vals) for lab, vals in read_store(store, buckets=[d], month=month)]
    return [reduce_days(lab, vals) for lab, vals in read_store(store, mp_ids=ids, month=month)]


def _fix_multimeter(days: pd.DataFrame, verbose: bool, exact_fn) -> pd.DataFrame:
    dup = days.duplicated(KEY, keep=False)
    if not dup.any():
        return days
    multi = set(days.loc[dup, "mp_id"].unique())
    if verbose:
        print(f"{len(multi):,} MP IDs have several meters -> summing their meters per day ...", flush=True)
    days = days[~days["mp_id"].isin(multi)]
    return pd.concat([days] + exact_fn(multi), ignore_index=True)


def _exact_days_from_store(store: str, mp_ids: Set[str], month: Optional[str]) -> pd.DataFrame:
    labs, vs = [], []
    for lab, vals in read_store(store, mp_ids=mp_ids, month=month):
        labs.append(lab); vs.append(vals)
    return _sum_meters(labs, vs)


def _sum_meters(labs, vs) -> pd.DataFrame:
    if not labs:
        return pd.DataFrame(columns=KEY + ["n_neg"])
    lab = pd.concat(labs, ignore_index=True)
    mat = pd.DataFrame(np.vstack(vs), columns=range(N_SLOTS))
    mat[KEY] = lab[KEY]
    summed = mat.groupby(KEY, sort=False)[list(range(N_SLOTS))].sum(min_count=1).reset_index()
    summed = summed.merge(lab.groupby(KEY, sort=False)[["n_neg", "vmax_raw"]]
                          .agg({"n_neg": "sum", "vmax_raw": "max"}).reset_index(), on=KEY)
    return reduce_days(summed[KEY + ["n_neg", "vmax_raw"]], summed[list(range(N_SLOTS))].to_numpy(dtype="float32"))


def _exact_days_for(path: str, mp_ids: Set[str], cache_dir: Optional[str] = None) -> pd.DataFrame:
    labs, vs = [], []
    for lab, vals in iter_aew_csv(path, mp_ids=mp_ids, cache_dir=cache_dir):
        labs.append(lab); vs.append(vals)
    return _sum_meters(labs, vs)


# =============================================================================
# 3. the twenty battery features, vectorised over all houses from the per-day table
# =============================================================================
def _season_or_all(I: pd.DataFrame, months: Tuple[int, ...], houses: pd.Index):
    """Season rows of one channel; a house with NO day in that season falls back to all its days."""
    sel = I[I["month"].isin(months)]
    n = sel.groupby("mp_id").size().reindex(houses).fillna(0).astype(int)
    missing = n.index[n == 0]
    if len(missing):
        sel = pd.concat([sel, I[I["mp_id"].isin(missing)]], ignore_index=True)
    return sel, n


def _log_flatness(a: pd.Series, b: pd.Series, eps: float) -> pd.Series:
    """-abs(log((a + eps) / (b + eps))): 0 = a and b are the same size (flat across the split),
    very negative = a and b differ a lot (strongly split, e.g. seasonal)."""
    ratio = (a + eps) / (b + eps)
    return -np.log(ratio).abs()


def feature_table(days: pd.DataFrame) -> pd.DataFrame:
    """One row per MP ID: the twenty battery features + data-quality columns."""
    days = days.copy()
    days["month"] = pd.to_datetime(days["date"]).dt.month
    I = days[days["channel"] == "import"]
    E = days[days["channel"] == "export"]
    houses = pd.Index(I["mp_id"].unique(), name="mp_id")
    out = pd.DataFrame(index=houses)
    g = lambda d, c: d.groupby("mp_id")[c]
    eps = ZERO_KW

    def R(x, fill=0.0) -> pd.Series:
        return x.reindex(houses).fillna(fill)

    Is, n_s = _season_or_all(I, SUMMER_MONTHS, houses)
    Iw, n_w = _season_or_all(I, WINTER_MONTHS, houses)
    Es, _ = _season_or_all(E, SUMMER_MONTHS, houses)
    Ew, _ = _season_or_all(E, WINTER_MONTHS, houses)

    imp_all = R(g(I, "tot").sum())
    n_all = R(g(I, "tot").size())
    n_s_used, n_w_used = R(g(Is, "tot").size()), R(g(Iw, "tot").size())
    mean_import_kw = R(g(I, "tot").sum() / g(I, "n_valid").sum().clip(lower=1))
    den_all = imp_all.clip(lower=eps)

    # ---------------- A. export counter ----------------
    exp_all = R(g(E, "tot").sum())
    out["export_kwh_total"] = exp_all / 4.0
    out["export_days_frac"] = R((E["tot"] / 4.0 > 0.1).groupby(E["mp_id"]).sum()) / n_all.clip(lower=1)
    out["export_offsolar_over_import"] = R(g(E, "night_tot").sum()) / den_all
    out["export_offsolar_days_frac"] = (R((E["night_tot"] / 4.0 > OFFSOLAR_KWH_DAY).groupby(E["mp_id"]).sum())
                                        / n_all.clip(lower=1))
    ok_plateau = E[E["plateau_frac"].notna()]
    out["export_flat_plateau_frac"] = R(g(ok_plateau, "plateau_frac").mean())
    out["export_bursts_per_day"] = R(g(E, "bursts").mean())
    exp_max = R(g(E, "vmax").quantile(0.95))
    out["export_max_kw"] = exp_max
    out["export_max_over_mean_import"] = exp_max / mean_import_kw.clip(lower=eps)
    exp_s_per_day = R(g(Es, "tot").sum()) / n_s_used.clip(lower=1) / 4.0
    exp_w_per_day = R(g(Ew, "tot").sum()) / n_w_used.clip(lower=1) / 4.0
    out["export_summer_winter_flatness"] = _log_flatness(exp_s_per_day, exp_w_per_day, eps)

    # ---------------- B. import counter ----------------
    ok = I[I["n_valid_night"] >= 0.9 * NIGHT_MASK.sum()]
    out["import_zero_hours_night"] = R(g(ok, "zero_q_night").mean()) / 4.0

    ok_c = I[I["charge_cnt"] >= 0.9 * CHARGE_MASK.sum()].copy()
    charge_mean = ok_c["charge_sum"] / ok_c["charge_cnt"].clip(lower=1)
    ok_c["charge_mean"] = charge_mean
    out["import_night_charge_kw_over_mean"] = R(g(ok_c, "charge_mean").quantile(0.95)) / mean_import_kw.clip(lower=eps)
    charge_var = ok_c["charge_sqsum"] / ok_c["charge_cnt"].clip(lower=1) - charge_mean ** 2
    ok_c["charge_cv"] = np.sqrt(charge_var.clip(lower=0)) / charge_mean.clip(lower=eps)
    out["import_night_charge_flatness"] = R(g(ok_c, "charge_cv").mean(), fill=1.0)
    thresh = (1.5 * mean_import_kw).reindex(ok_c["mp_id"]).to_numpy()
    ok_c["above_1p5x"] = (ok_c["charge_mean"].to_numpy() > thresh).astype(float)
    out["import_night_charge_persistence"] = R(g(ok_c, "above_1p5x").mean())

    ok_p = I[(I["n_valid_peak"] >= 0.9 * PEAK_MASK.sum()) & (I["n_valid"] - I["n_valid_peak"] > 0)].copy()
    peak_mean = ok_p["peak_tot"] / ok_p["n_valid_peak"].clip(lower=1)
    rest_mean = (ok_p["tot"] - ok_p["peak_tot"]) / (ok_p["n_valid"] - ok_p["n_valid_peak"]).clip(lower=1)
    ok_p["suppression"] = 1.0 - peak_mean / rest_mean.clip(lower=eps)
    out["import_evening_peak_suppression"] = R(g(ok_p, "suppression").mean())

    out["import_ramp_events_night"] = R(g(I, "ramp_events_night").mean())

    ok_d = I[I["n_valid_day"] >= 0.9 * DAY_MASK.sum()]
    out["import_zero_hours_midday_allyear"] = R(g(ok_d, "zero_q_day").mean()) / 4.0

    imp_s_per_day = R(g(Is, "tot").sum()) / n_s_used.clip(lower=1)
    imp_w_per_day = R(g(Iw, "tot").sum()) / n_w_used.clip(lower=1)
    out["import_seasonal_flatness"] = _log_flatness(imp_s_per_day, imp_w_per_day, eps)

    ok_m = I[(I["n_valid_mid"] >= 0.9 * MID_MASK.sum()) & (I["n_valid_night"] >= 0.9 * NIGHT_MASK.sum())].copy()
    ok_m["mid_mean"] = ok_m["mid_tot"] / ok_m["n_valid_mid"].clip(lower=1)
    ok_m["night_mean"] = ok_m["night_tot"] / ok_m["n_valid_night"].clip(lower=1)
    out["import_midday_night_ratio_allyear"] = (R(g(ok_m, "mid_mean").mean())
                                                / R(g(ok_m, "night_mean").mean()).clip(lower=eps))

    # ---------------- C. joint ----------------
    out["daily_activity_cycles"] = out["export_bursts_per_day"] + R(g(I, "charge_bursts").mean())

    merged = E[["mp_id", "date", "mid_tot"]].merge(I[["mp_id", "date", "peak_tot"]], on=["mp_id", "date"], how="inner")
    if len(merged):
        merged["xx"] = merged["mid_tot"] ** 2
        merged["yy"] = merged["peak_tot"] ** 2
        merged["xy"] = merged["mid_tot"] * merged["peak_tot"]
        mg = merged.groupby("mp_id")
        n_pair = mg.size()
        sx, sy = mg["mid_tot"].sum(), mg["peak_tot"].sum()
        sxx, syy, sxy = mg["xx"].sum(), mg["yy"].sum(), mg["xy"].sum()
        nn = n_pair.clip(lower=1)
        cov = sxy / nn - (sx / nn) * (sy / nn)
        varx = (sxx / nn - (sx / nn) ** 2).clip(lower=0)
        vary = (syy / nn - (sy / nn) ** 2).clip(lower=0)
        denom = np.sqrt(varx * vary)
        corr = pd.Series(np.where((n_pair >= 10) & (denom > 0), cov / denom.replace(0, np.nan), np.nan),
                         index=n_pair.index)
        out["self_consumption_shift_score"] = R(-corr)
    else:
        out["self_consumption_shift_score"] = 0.0

    # ---------------- data quality ----------------
    out["n_days"] = n_all.astype(int)
    out["n_summer_days"] = n_s.astype(int)
    out["n_winter_days"] = n_w.astype(int)
    out["has_export_register"] = houses.isin(E["mp_id"].unique()).astype(int)
    out["has_real_export"] = (exp_max >= EXPORT_REAL_KW).astype(int)
    out["has_offsolar_export"] = (out["export_offsolar_days_frac"] > 0.02).astype(int)
    out["export_max_raw_kw"] = R(g(E, "vmax_raw").max())
    out["mean_import_kw"] = mean_import_kw
    out["n_negative_raw"] = R(days.groupby("mp_id")["n_neg"].sum()).astype(int)
    return out


FEATURE_DESC = {
    "export_kwh_total": "total energy exported [kWh]; >0 = something behind the meter can push energy out",
    "export_days_frac": "share of days with >0.1 kWh exported",
    "export_offsolar_over_import": "export in the night window / import, all days -- the sun cannot do this; "
                                   "direct battery / storage signal",
    "export_offsolar_days_frac": "share of days with real export in that same night window (persistence, not just size)",
    "export_flat_plateau_frac": "of exporting quarter-hours, share within 10% of the day's max "
                                "(steady inverter plateau vs PV's smooth bell)",
    "export_bursts_per_day": "mean number of separate export episodes per day (PV alone ~1; cycling battery 2+)",
    "export_max_kw": "95th percentile of the daily max export [kW] ~ inverter size (0 if none)",
    "export_max_over_mean_import": "that max / the house's mean import power: generator / inverter size relative to the house",
    "export_summer_winter_flatness": "-|log(summer export/day / winter export/day)|; near 0 = flat across seasons "
                                     "(battery-like); very negative = seasonal (PV-like). Caveat: 0 can also mean "
                                     "no export at all -- check has_export_register.",
    "import_zero_hours_night": "hours/day (22h-06h) with import <= 20 W, all year (the sun cannot cause this)",
    "import_night_charge_kw_over_mean": "95th percentile of mean import power in 23h-06h / overall mean import power "
                                        "(high and sustained -> scheduled grid charging; also matches a home EV charger)",
    "import_night_charge_flatness": "coefficient of variation of import power in 23h-06h on well-covered nights "
                                    "(lower = steadier, charger-like)",
    "import_night_charge_persistence": "share of nights on which the mean 23h-06h import exceeds 1.5x the house's "
                                       "overall mean import",
    "import_evening_peak_suppression": "1 - (mean import 17h-21h) / (mean import rest of day): positive and large -> "
                                       "evening peak shaved, consistent with discharge",
    "import_ramp_events_night": "mean count/day of import jumps > 0.3 kW between consecutive quarter-hours within 22h-06h",
    "import_zero_hours_midday_allyear": "hours/day (9h-17h) with import <= 20 W, ALL YEAR (persistence in winter "
                                        "is not explained by PV alone)",
    "import_seasonal_flatness": "-|log(summer daily import / winter daily import)|; near 0 = small seasonal shape "
                                "difference (buffered load). Same 0-can-mean-no-data caveat as above.",
    "import_midday_night_ratio_allyear": "mean import 11h-15h / max(mean import 22h-06h, 20 W), all year",
    "daily_activity_cycles": "export bursts/day + distinct 23h-06h charging episodes/day: overall system cycling rate",
    "self_consumption_shift_score": "-corr(midday export, same-day evening-peak import) across days; positive = "
                                    "time-shift signature (store at midday, use in the evening)",
    "n_days": "days of data", "n_summer_days": "May-Aug days (0 = summer features used all days)",
    "n_winter_days": "Nov-Feb days (0 = winter features used all days)",
    "has_export_register": "1 = meter has export rows",
    "has_real_export": "1 = typical daily max export >= 0.1 kW",
    "has_offsolar_export": "1 = real export in the night window on a non-trivial share of days",
    "export_max_raw_kw": "raw max export before the 50 W noise floor",
    "mean_import_kw": "average import power [kW]: household 0.2-0.8, tens = business",
    "n_negative_raw": "negative raw values floored (should be 0)",
}

# Value used when a feature cannot be computed.  Chosen as the value a house WITHOUT a battery would
# show, so a fill never looks like a detection on its own -- but see the module docstring caveats for
# export_summer_winter_flatness / import_seasonal_flatness, where a neutral-looking 0 can also come
# from a genuine absence of data, not from evidence of flatness.
FILL_VALUES = {
    "export_kwh_total": 0.0, "export_days_frac": 0.0, "export_offsolar_over_import": 0.0,
    "export_offsolar_days_frac": 0.0, "export_flat_plateau_frac": 0.0, "export_bursts_per_day": 0.0,
    "export_max_kw": 0.0, "export_max_over_mean_import": 0.0, "export_summer_winter_flatness": 0.0,
    "import_zero_hours_night": 0.0, "import_night_charge_kw_over_mean": 0.0, "import_night_charge_flatness": 1.0,
    "import_night_charge_persistence": 0.0,
    "import_evening_peak_suppression": 0.0, "import_ramp_events_night": 0.0,
    "import_zero_hours_midday_allyear": 0.0, "import_seasonal_flatness": 0.0, "import_midday_night_ratio_allyear": 1.0,
    "daily_activity_cycles": 0.0, "self_consumption_shift_score": 0.0,
    "has_real_export": 0, "export_max_raw_kw": 0.0, "mean_import_kw": 0.0,
}


def fill_missing(tab: pd.DataFrame) -> pd.DataFrame:
    tab = tab.copy()
    for c, v in FILL_VALUES.items():
        if c in tab.columns:
            tab[c] = tab[c].fillna(v)
    tab["has_real_export"] = tab["has_real_export"].astype(int)
    return tab


def write_xlsx(tab: pd.DataFrame, path: str) -> bool:
    """Excel copy of the feature table: houses as rows, features as columns, frozen header,
    filters, colour scales per feature, plus a 'legend' sheet."""
    try:
        import openpyxl
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("(no Excel copy: pip install openpyxl)", flush=True)
        return False
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        tab.reset_index().to_excel(xw, sheet_name="features", index=False, float_format="%.4f")
        pd.DataFrame({"feature": list(FEATURE_DESC), "meaning": list(FEATURE_DESC.values())}
                     ).to_excel(xw, sheet_name="legend", index=False)
        ws = xw.sheets["features"]
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        n_rows = len(tab) + 1
        for j, col in enumerate(tab.reset_index().columns, start=1):
            letter = get_column_letter(j)
            ws.column_dimensions[letter].width = max(12, min(34, len(str(col)) * 0.9))
            if col in FEATURE_NAMES and n_rows > 2:
                ws.conditional_formatting.add(f"{letter}2:{letter}{n_rows}",
                    ColorScaleRule(start_type="min", start_color="F7FBFF", end_type="max", end_color="2171B5"))
        for cell in ws[1]:
            cell.alignment = openpyxl.styles.Alignment(text_rotation=60, vertical="bottom")
        ws.row_dimensions[1].height = 150
        lg = xw.sheets["legend"]
        lg.column_dimensions["A"].width = 42; lg.column_dimensions["B"].width = 100
    return True


# =============================================================================
# 4. single-house helpers (mirror pv20_inspect.py's load_house / house_features)
# =============================================================================
def load_house(root: str, mp_id: str, month: Optional[str] = None,
               max_files: Optional[int] = None) -> Dict[str, pd.DataFrame]:
    labs, vs = [], []
    for f in find_files(root, month, max_files):
        for lab, vals in iter_aew_csv(f, mp_ids=[str(mp_id)]):
            labs.append(lab); vs.append(vals)
    if not labs:
        raise SystemExit(f"MP ID {mp_id} not found in the selected files")
    lab = pd.concat(labs, ignore_index=True)
    mat = pd.DataFrame(np.vstack(vs), columns=range(N_SLOTS))
    mat[["channel", "date"]] = lab[["channel", "date"]]
    house: Dict[str, pd.DataFrame] = {"n_negative_raw": int(lab["n_neg"].sum())}
    for ch, gch in mat.groupby("channel"):
        house[ch] = gch.groupby("date")[list(range(N_SLOTS))].sum(min_count=1).sort_index()
    return house


def house_features(house: Dict[str, pd.DataFrame]) -> Dict[str, float]:
    parts = []
    for ch in ("import", "export"):
        if ch in house:
            m = house[ch]
            lab = pd.DataFrame({"mp_id": "x", "channel": ch, "date": pd.to_datetime(m.index), "n_neg": 0})
            parts.append(reduce_days(lab, m.to_numpy(dtype="float32")))
    days = pd.concat(parts, ignore_index=True)
    days.loc[days.index[0], "n_neg"] = int(house.get("n_negative_raw", 0))
    return feature_table(days).iloc[0].to_dict()


# =============================================================================
# 5. command line
# =============================================================================
def main(argv: Optional[List[str]] = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("root", nargs="?", help="folder containing the monthly AEW CSVs (not needed with a store)")
    ap.add_argument("--store", help=f"binary store built by build_store.py (default: ./{STORE_DEFAULT} if it exists)")
    ap.add_argument("-o", "--out", default="battery20_features.csv")
    ap.add_argument("--month", help='only files whose path contains this text, e.g. "Juni 2023" or "2023"')
    ap.add_argument("--max-files", type=int, help="read only the first N monthly files")
    ap.add_argument("--max-houses", type=int, help="the first N MP IDs of the first file, followed through ALL files")
    ap.add_argument("--mp-ids", help="comma-separated MP IDs to keep, e.g. 53628,53701")
    ap.add_argument("--mp-ids-file", help="text file with one MP ID per line (e.g. all labelled houses)")
    ap.add_argument("--pattern", default=FILE_GLOB, help=f"file name pattern of the meter files (default {FILE_GLOB})")
    ap.add_argument("--keep-nan", action="store_true",
                    help="leave non-computable features empty instead of the neutral fill values")
    ap.add_argument("--xlsx", help="Excel copy path (default: same name as -o with .xlsx)")
    ap.add_argument("--no-xlsx", action="store_true", help="do not write the Excel copy")
    ap.add_argument("--refresh-days", action="store_true",
                    help="store: recompute the all-houses per-day table instead of using _battery_days.parquet")
    ap.add_argument("--workers", type=int, default=4,
                    help="files skimmed in parallel when a house list is given (default 4)")
    ap.add_argument("--cache-dir", default="aew_cache",
                    help="folder for the extracted lines (default ./aew_cache; '' = off)")
    a = ap.parse_args(argv)

    store = a.store or (STORE_DEFAULT if os.path.isdir(STORE_DEFAULT) else None)
    if store and a.root and a.store is None:
        print(f"note: using the store ./{STORE_DEFAULT} instead of the CSVs (pass --store '' to force the CSVs)", flush=True)
    if a.store == "":
        store = None
    ids = [x.strip() for x in a.mp_ids.split(",")] if a.mp_ids else None
    if a.mp_ids_file:
        with open(a.mp_ids_file) as f:
            ids = (ids or []) + [ln.strip().split(";")[0].split(",")[0] for ln in f
                                 if ln.strip() and not ln.lower().startswith("mp")]
    days = load_days(a.root, month=a.month, max_files=a.max_files, max_houses=a.max_houses, mp_ids=ids,
                     pattern=a.pattern, workers=a.workers, cache_dir=a.cache_dir or None, store=store,
                     use_days_cache=not a.refresh_days)
    tab = feature_table(days)
    n_missing = int(tab[FEATURE_NAMES].isna().sum().sum())
    if n_missing:
        print(f"warning: {n_missing} feature values could not be computed" + ("" if a.keep_nan else " -> filled"), flush=True)
    if not a.keep_nan:
        tab = fill_missing(tab)
    tab.to_csv(a.out, float_format="%.4f")
    xlsx = a.xlsx if a.xlsx else os.path.splitext(a.out)[0] + ".xlsx"
    ok = False if a.no_xlsx else write_xlsx(tab, xlsx)
    print(f"\n{len(tab):,} houses -> {a.out}" + (f"  +  {xlsx}" if ok else "") + "\n")
    with pd.option_context("display.width", 250, "display.max_columns", 40, "display.max_rows", 60,
                           "display.float_format", "{:.3f}".format):
        if len(tab) <= 30:
            print(tab.T)
        else:
            print(tab[FEATURE_NAMES].describe().T[["count", "mean", "min", "25%", "50%", "75%", "max"]])
    n_nan = tab[FEATURE_NAMES].isna().all()
    if n_nan.any():
        print("\nall-NaN features (no summer / winter days in the selected files, or no export rows):",
              ", ".join(n_nan[n_nan].index))


if __name__ == "__main__":
    main()